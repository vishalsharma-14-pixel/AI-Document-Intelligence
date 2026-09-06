import pdfplumber
from docx import Document as DocxDocument
from openpyxl import load_workbook


class ParsingError(Exception):
    pass


def parse_document(file_path: str, file_type: str) -> tuple[list[dict], int | None]:
    """Extract text units from a document.

    Returns (units, page_count) where each unit is
    {"text": str, "page": int | None, "section": str | None}.
    """
    if file_type == "pdf":
        return _parse_pdf(file_path)
    if file_type == "docx":
        return _parse_docx(file_path), None
    if file_type in ("xlsx", "xls"):
        return _parse_xlsx(file_path), None
    if file_type in ("txt", "md"):
        return _parse_text(file_path), None
    raise ParsingError(f"Unsupported file type: {file_type}")


def _parse_pdf(file_path: str) -> tuple[list[dict], int]:
    units: list[dict] = []
    try:
        with pdfplumber.open(file_path) as pdf:
            page_count = len(pdf.pages)
            for i, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    units.append({"text": text, "page": i, "section": None})
    except Exception as exc:
        raise ParsingError(f"Failed to parse PDF: {exc}") from exc
    if not units:
        raise ParsingError("No extractable text found in PDF (it may be scanned/image-only)")
    return units, page_count


def _parse_docx(file_path: str) -> list[dict]:
    units: list[dict] = []
    try:
        doc = DocxDocument(file_path)
        current_section = "Document"
        buffer: list[str] = []

        def flush():
            text = "\n".join(buffer).strip()
            if text:
                units.append({"text": text, "page": None, "section": current_section})

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            is_heading = para.style is not None and para.style.name.lower().startswith("heading")
            if is_heading:
                flush()
                current_section = text
                buffer = []
            else:
                buffer.append(text)
        flush()
    except Exception as exc:
        raise ParsingError(f"Failed to parse DOCX: {exc}") from exc
    if not units:
        raise ParsingError("No extractable text found in DOCX")
    return units


def _parse_xlsx(file_path: str) -> list[dict]:
    units: list[dict] = []
    try:
        wb = load_workbook(file_path, data_only=True, read_only=True)
        for sheet in wb.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
            text = "\n".join(lines).strip()
            if text:
                units.append({"text": text, "page": None, "section": sheet.title})
    except Exception as exc:
        raise ParsingError(f"Failed to parse XLSX: {exc}") from exc
    if not units:
        raise ParsingError("No extractable data found in spreadsheet")
    return units


def _parse_text(file_path: str) -> list[dict]:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
    except Exception as exc:
        raise ParsingError(f"Failed to read text file: {exc}") from exc
    if not text:
        raise ParsingError("File is empty")
    return [{"text": text, "page": None, "section": None}]
